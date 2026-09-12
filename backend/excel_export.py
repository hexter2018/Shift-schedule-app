"""
Excel generation — Python/openpyxl port of frontend/src/lib/excel.js.

Why this exists as a second implementation instead of reusing the JS one:
the approval workflow needs the backend to (a) generate the file at submit
time and (b) re-open that *exact same file* later to inject signatures,
without a browser in the loop. openpyxl is the natural tool for that on
the Python side.

Keep this in sync with excel.js if the visual template changes — column
widths, fonts, tint formula, and the group-separator positions are
intentionally identical to that file so a schedule looks the same whether
downloaded directly or produced via the approval flow.
"""

from dataclasses import dataclass, asdict
from datetime import date, datetime
import calendar

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

FONT_NAME = "TH Sarabun New"
FONT_SIZE = 18
THAI_MONTHS = ["มกราคม","กุมภาพันธ์","มีนาคม","เมษายน","พฤษภาคม","มิถุนายน",
               "กรกฎาคม","สิงหาคม","กันยายน","ตุลาคม","พฤศจิกายน","ธันวาคม"]
THAI_WEEKDAYS = ["จ","อ","พ","พฤ","ศ","ส","อา"]  # Python weekday(): 0=Mon
GRID = "D8D2C4"
# Blank rows inserted after these 1-based employee positions — keep this in
# sync with GROUP_SEPARATOR_AFTER in excel.js.
GROUP_SEPARATOR_AFTER = {4, 9, 13, 17}


def days_in_month(year_be: int, month: int) -> int:
    return calendar.monthrange(year_be - 543, month)[1]


def weekday_label(year_be: int, month: int, day: int) -> str:
    return THAI_WEEKDAYS[date(year_be - 543, month, day).weekday()]


def is_weekend(year_be: int, month: int, day: int) -> bool:
    return date(year_be - 543, month, day).weekday() >= 5


def is_holiday(year_be: int, month: int, day: int, holidays: list) -> dict | None:
    for h in holidays:
        if h["month"] == month and h["day"] == day and (h.get("year") is None or h["year"] == year_be):
            return h
    return None


def parse_cell_value(raw: str) -> dict:
    if not raw:
        return {"base": "", "wh": False, "th": False, "ot": False}
    v = raw.strip().upper()
    for suffix, key in (("WH", "wh"), ("TH", "th"), ("OT", "ot")):
        if v.endswith(suffix) and len(v) > 2:
            return {"base": v[:-2], "wh": key == "wh", "th": key == "th", "ot": key == "ot"}
    return {"base": v, "wh": False, "th": False, "ot": False}


def tint_hex(hex_color: str, amount: float) -> str:
    h = hex_color.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    mix = lambda c: round(c + (255 - c) * amount)
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def argb(hex_color: str) -> str:
    """openpyxl wants 8-hex ARGB; without the FF alpha prefix colors render
    as fully transparent (bit us once already on the JS side)."""
    return "FF" + hex_color.lstrip("#").upper()


@dataclass
class SignatureLayout:
    """Recorded at generation time and stored on the approval_cycles row —
    the callback reads these back instead of assuming fixed coordinates,
    since they shift with employee count / holidays that month."""
    sig_line_row: int   # blank underline row -> gets the approver's name
    sig_label_row: int  # static "ผู้จัดการแผนก" / "ผู้จัดการส่วน" caption -> untouched
    sig_date_row: int   # "……./……./……" placeholder -> gets the approval date
    sig_cert_row: int   # blank until approval -> gets the Power Automate
                         # Approval ID as a small audit-trail reference,
                         # e.g. "รหัสอ้างอิง: 08584125..."
    block1_start: int   # Section Manager's column (1-indexed)
    block2_start: int   # Division Manager's column (1-indexed)
    block_width: int


def build_workbook(state: dict, holidays: list) -> tuple[Workbook, SignatureLayout]:
    nd = days_in_month(state["yearBE"], state["month"])
    NUM_FIXED_COLS = 3
    last_col = NUM_FIXED_COLS + nd + 1  # 1-indexed: note column

    wb = Workbook()
    ws = wb.active
    ws.title = "ตารางกะ"

    base_font = Font(name=FONT_NAME, size=FONT_SIZE)
    bold_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
    grid_side = Side(style="thin", color=argb(GRID))
    all_sides = Border(top=grid_side, bottom=grid_side, left=grid_side, right=grid_side)
    center = Alignment(horizontal="center", vertical="center")
    center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    def set_row_height(row, pt):
        ws.row_dimensions[row].height = pt

    # ---- title block ----
    ws.cell(row=1, column=1, value="ตารางปฏิบัติงานประจำเดือน")
    ws.cell(row=2, column=1, value=state.get("department", ""))
    ws.cell(row=3, column=1, value=f"ประจำเดือน {THAI_MONTHS[state['month']-1]} {state['yearBE']}")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws.cell(row=1, column=1).font = bold_font
    ws.cell(row=2, column=1).font = bold_font
    ws.cell(row=3, column=1).font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=argb("6B645C"))
    for r in (1, 2, 3):
        ws.cell(row=r, column=1).alignment = center
    set_row_height(1, 27); set_row_height(2, 27); set_row_height(3, 27)
    set_row_height(4, 6)  # spacer

    # ---- header rows (5 = labels/day numbers, 6 = weekday letters) ----
    ws.cell(row=5, column=1, value="ลำดับ")
    ws.cell(row=5, column=2, value="รหัส")
    ws.cell(row=5, column=3, value="ชื่อ-สกุล")
    ws.cell(row=5, column=last_col, value="หมายเหตุ")
    ws.merge_cells(start_row=5, start_column=1, end_row=6, end_column=1)
    ws.merge_cells(start_row=5, start_column=2, end_row=6, end_column=2)
    ws.merge_cells(start_row=5, start_column=3, end_row=6, end_column=3)
    ws.merge_cells(start_row=5, start_column=last_col, end_row=6, end_column=last_col)
    set_row_height(5, 24); set_row_height(6, 18)

    for d in range(1, nd + 1):
        col = NUM_FIXED_COLS + d
        hol = is_holiday(state["yearBE"], state["month"], d, holidays)
        weekend = is_weekend(state["yearBE"], state["month"], d)
        fill_hex = "FDEAEA" if hol else ("FAF6EE" if weekend else "FAF8F4")
        wd_color = "B91C1C" if hol else "B45309"
        ws.cell(row=5, column=col, value=d)
        ws.cell(row=6, column=col, value=weekday_label(state["yearBE"], state["month"], d))
        for r in (5, 6):
            c = ws.cell(row=r, column=col)
            c.fill = PatternFill("solid", fgColor=argb(fill_hex))
            c.border = all_sides
            c.alignment = center_wrap if r == 5 else center
        ws.cell(row=5, column=col).font = bold_font
        ws.cell(row=6, column=col).font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=argb(wd_color))

    for c in (1, 2, 3, last_col):
        for r in (5, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = bold_font
            cell.fill = PatternFill("solid", fgColor=argb("FAF8F4"))
            cell.border = all_sides
            cell.alignment = center_wrap if r == 5 else center

    # ---- employee data rows (with group-separator blank rows) ----
    shift_codes = {c["code"].upper(): c for c in state["shiftCodes"]}
    row = 7
    for idx, emp in enumerate(state["employees"]):
        ws.cell(row=row, column=1, value=idx + 1).font = base_font
        ws.cell(row=row, column=2, value=emp.get("empCode", "")).font = base_font
        name_cell = ws.cell(row=row, column=3, value=emp.get("name", ""))
        name_cell.font = bold_font
        name_cell.alignment = left
        for c in (1, 2):
            ws.cell(row=row, column=c).alignment = center
        for d in range(1, nd + 1):
            col = NUM_FIXED_COLS + d
            raw = (emp.get("days", {}).get(str(d)) or emp.get("days", {}).get(d) or "")
            parsed = parse_cell_value(raw)
            base = parsed["base"]
            cell = ws.cell(row=row, column=col, value=raw)
            cell.alignment = center
            defn = shift_codes.get(base)
            if defn:
                tint_amt = 0.55 if (parsed["wh"] or parsed["th"]) else 0.86
                bg = tint_hex(defn["color"], tint_amt)
                fg = defn["color"].lstrip("#").upper()
                cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=argb(fg))
                cell.fill = PatternFill("solid", fgColor=argb(bg))
                if parsed["wh"]:
                    s = Side(style="dashed", color=argb(fg)); cell.border = Border(top=s, bottom=s, left=s, right=s)
                elif parsed["th"]:
                    s = Side(style="dotted", color=argb("B91C1C")); cell.border = Border(top=s, bottom=s, left=s, right=s)
                elif parsed["ot"]:
                    s = Side(style="medium", color=argb("DB2777")); cell.border = Border(top=s, bottom=s, left=s, right=s)
                else:
                    cell.border = all_sides
            elif base == "WH":
                cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color=argb("6B645C"))
                cell.fill = PatternFill("solid", fgColor=argb("EFECE4"))
                cell.border = all_sides
            else:
                cell.font = base_font
                cell.border = all_sides
                hol = is_holiday(state["yearBE"], state["month"], d, holidays)
                if hol:
                    cell.fill = PatternFill("solid", fgColor=argb("FDEAEA"))
        for c in (1, 2, 3):
            ws.cell(row=row, column=c).border = all_sides
        note_cell = ws.cell(row=row, column=last_col, value=emp.get("note", ""))
        note_cell.font = base_font
        note_cell.alignment = left
        note_cell.border = all_sides
        set_row_height(row, 27)
        row += 1

        if (idx + 1) in GROUP_SEPARATOR_AFTER:
            for c in range(1, last_col + 1):
                cell = ws.cell(row=row, column=c)
                cell.font = base_font
                cell.border = all_sides
            set_row_height(row, 27)
            row += 1

    # ---- holiday list (only if this month has any) ----
    month_holidays = [
        {"d": d, "name": h["name"]}
        for d in range(1, nd + 1)
        if (h := is_holiday(state["yearBE"], state["month"], d, holidays))
    ]
    if month_holidays:
        row += 1
        ws.cell(row=row, column=1, value="วันหยุดในเดือนนี้").font = Font(
            name=FONT_NAME, size=FONT_SIZE, bold=True, color=argb("B91C1C"))
        set_row_height(row, 24)
        row += 1
        for h in month_holidays:
            ws.cell(row=row, column=1, value=f"วันที่ {h['d']}").font = Font(
                name=FONT_NAME, size=FONT_SIZE, bold=True, color=argb("B91C1C"))
            ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=argb("FDEAEA"))
            ws.cell(row=row, column=1).alignment = center
            ws.cell(row=row, column=2, value=h["name"]).font = base_font
            set_row_height(row, 24)
            row += 1

    # ---- signature block: 4 blank spacer rows, then line / label / date ----
    row += 4
    sig_block_width = 4
    sig_gap = 2
    block2_end = last_col - 1  # last *day* column, not the note column
    block2_start = max(1, block2_end - sig_block_width + 1)
    block1_end = max(1, block2_start - sig_gap - 1)
    block1_start = max(1, block1_end - sig_block_width + 1)

    sig_line_row = row
    ws.merge_cells(start_row=sig_line_row, start_column=block1_start,
                    end_row=sig_line_row, end_column=block1_start + sig_block_width - 1)
    ws.merge_cells(start_row=sig_line_row, start_column=block2_start,
                    end_row=sig_line_row, end_column=block2_start + sig_block_width - 1)
    dash = Side(style="dashed", color=argb("6B645C"))
    for c in (block1_start, block2_start):
        ws.cell(row=sig_line_row, column=c).border = Border(bottom=dash)
        ws.cell(row=sig_line_row, column=c).alignment = center
    set_row_height(sig_line_row, 26.1)
    row += 1

    sig_label_row = row
    ws.cell(row=sig_label_row, column=block1_start, value="ผู้จัดการแผนก")
    ws.cell(row=sig_label_row, column=block2_start, value="ผู้จัดการส่วน")
    ws.merge_cells(start_row=sig_label_row, start_column=block1_start,
                    end_row=sig_label_row, end_column=block1_start + sig_block_width - 1)
    ws.merge_cells(start_row=sig_label_row, start_column=block2_start,
                    end_row=sig_label_row, end_column=block2_start + sig_block_width - 1)
    for c in (block1_start, block2_start):
        cell = ws.cell(row=sig_label_row, column=c)
        cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color=argb("6B645C"))
        cell.alignment = center
    set_row_height(sig_label_row, 27)
    row += 1

    sig_date_row = row
    ws.cell(row=sig_date_row, column=block1_start, value="……./……./……")
    ws.cell(row=sig_date_row, column=block2_start, value="……./……./……")
    ws.merge_cells(start_row=sig_date_row, start_column=block1_start,
                    end_row=sig_date_row, end_column=block1_start + sig_block_width - 1)
    ws.merge_cells(start_row=sig_date_row, start_column=block2_start,
                    end_row=sig_date_row, end_column=block2_start + sig_block_width - 1)
    for c in (block1_start, block2_start):
        cell = ws.cell(row=sig_date_row, column=c)
        cell.font = base_font
        cell.alignment = center
    row += 1

    # Reserved, left blank until the outcome callback injects the actual
    # Power Automate Approval ID — small/muted so it reads as a reference
    # number, not competing visually with the name/date above it.
    sig_cert_row = row
    ws.merge_cells(start_row=sig_cert_row, start_column=block1_start,
                    end_row=sig_cert_row, end_column=block1_start + sig_block_width - 1)
    ws.merge_cells(start_row=sig_cert_row, start_column=block2_start,
                    end_row=sig_cert_row, end_column=block2_start + sig_block_width - 1)
    cert_font = Font(name=FONT_NAME, size=11, color=argb("94A3B8"))
    for c in (block1_start, block2_start):
        cell = ws.cell(row=sig_cert_row, column=c)
        cell.font = cert_font
        cell.alignment = center
    set_row_height(sig_cert_row, 15)

    # ---- column widths ----
    ws.column_dimensions["A"].width = 6.02
    ws.column_dimensions["B"].width = 12.02
    ws.column_dimensions["C"].width = 33.16
    for d in range(1, nd + 1):
        ws.column_dimensions[get_column_letter(NUM_FIXED_COLS + d)].width = 9.02
    ws.column_dimensions[get_column_letter(last_col)].width = 47.02
    ws.page_setup.orientation = "landscape"

    layout = SignatureLayout(
        sig_line_row=sig_line_row, sig_label_row=sig_label_row, sig_date_row=sig_date_row,
        sig_cert_row=sig_cert_row,
        block1_start=block1_start, block2_start=block2_start, block_width=sig_block_width,
    )
    return wb, layout


def inject_approval_signatures(
    pending_path: str,
    layout: SignatureLayout,
    section_name: str, section_approved_at: str,
    division_name: str, division_approved_at: str,
    section_approval_id: str | None = None,
    division_approval_id: str | None = None,
) -> Workbook:
    """Re-opens the pending file and writes into the anchor (top-left) cell
    of each merged signature range — openpyxl only accepts writes there;
    writing to any other cell in a merged range raises AttributeError.

    section_approval_id / division_approval_id are Power Automate's own
    "Approval ID" for each "Start and wait for an approval" step — a
    unique, Microsoft-assigned identifier for that specific approval
    instance. Including it on the signed document gives a verifiable
    audit-trail reference distinct from just a name and a timestamp: it
    ties the printed signature back to the exact approval record in
    Power Automate's own run history, which a name/date pair alone
    doesn't. Optional (defaults to None) so this function's older
    call signature — before cert IDs existed — still works; callers that
    don't have an ID just leave that block's reference line blank."""
    wb = load_workbook(pending_path)
    ws = wb.active
    base_font = Font(name=FONT_NAME, size=FONT_SIZE)
    center = Alignment(horizontal="center", vertical="center")

    def fmt_thai_date(iso_str: str) -> str:
        dt = datetime.fromisoformat(iso_str.replace("Z", "+00:00"))
        return f"{dt.day:02d}/{dt.month:02d}/{dt.year + 543}"

    for col, name, when, cert_id in (
        (layout.block1_start, section_name, section_approved_at, section_approval_id),
        (layout.block2_start, division_name, division_approved_at, division_approval_id),
    ):
        name_cell = ws.cell(row=layout.sig_line_row, column=col, value=name)
        name_cell.font = base_font
        name_cell.alignment = center
        date_cell = ws.cell(row=layout.sig_date_row, column=col, value=fmt_thai_date(when))
        date_cell.font = base_font
        date_cell.alignment = center
        if cert_id:
            cert_cell = ws.cell(row=layout.sig_cert_row, column=col, value=f"รหัสอ้างอิง: {cert_id}")
            cert_cell.font = Font(name=FONT_NAME, size=11, color=argb("94A3B8"))
            cert_cell.alignment = center

    return wb
